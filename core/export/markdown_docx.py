"""Convert report Markdown into .md / .docx with citations and pending highlights."""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
import re
from typing import Literal

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_CITATION_RE = re.compile(r"\[([^\[\]]+?)\s+p(\d+)\]")
_PENDING = "【待核实】"


@dataclass
class ParsedBlock:
    kind: Literal["heading", "paragraph"]
    text: str
    level: int = 0


@dataclass
class ParsedReport:
    title: str
    blocks: list[ParsedBlock] = field(default_factory=list)
    citations: list[tuple[str, int]] = field(default_factory=list)


def _sanitize_filename(title: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", (title or "report").strip()) or "report"
    return cleaned[:80]


def parse_report_markdown(report_md: str, *, title: str) -> ParsedReport:
    lines = (report_md or "").splitlines()
    blocks: list[ParsedBlock] = []
    citations: list[tuple[str, int]] = []
    seen: set[tuple[str, int]] = set()
    buf: list[str] = []
    doc_title = title

    def flush() -> None:
        nonlocal buf
        text = "\n".join(buf).strip()
        buf = []
        if text:
            blocks.append(ParsedBlock(kind="paragraph", text=text))

    for raw in lines:
        heading = _HEADING_RE.match(raw.strip())
        if heading:
            flush()
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1 and not any(b.kind == "heading" for b in blocks):
                doc_title = text or doc_title
            blocks.append(ParsedBlock(kind="heading", text=text, level=level))
            continue
        if not raw.strip():
            flush()
            continue
        buf.append(raw.rstrip())
    flush()

    for block in blocks:
        for match in _CITATION_RE.finditer(block.text):
            key = (match.group(1).strip(), int(match.group(2)))
            if key not in seen:
                seen.add(key)
                citations.append(key)
    return ParsedReport(title=doc_title or title or "报告", blocks=blocks, citations=citations)


def default_output_path(title: str, fmt: Literal["docx", "md"], *, out_dir: Path | None = None) -> Path:
    root = out_dir if out_dir is not None else Path("output")
    root.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return root / f"{_sanitize_filename(title)}-{stamp}.{fmt}"


def write_markdown(report_md: str, *, title: str, path: Path | None = None) -> Path:
    target = path or default_output_path(title, "md")
    target.parent.mkdir(parents=True, exist_ok=True)
    body = report_md if (report_md or "").lstrip().startswith("#") else f"# {title}\n\n{report_md}"
    target.write_text(body if body.endswith("\n") else body + "\n", encoding="utf-8")
    return target


def _add_toc_paragraph(document: Document, headings: Sequence[str]) -> None:
    document.add_heading("目录", level=1)
    for index, name in enumerate(headings, start=1):
        document.add_paragraph(f"{index}. {name}")


def _append_runs_with_citations_and_pending(paragraph, text: str) -> None:
    """Split text into runs; yellow-highlight 【待核实】; keep [file pN] as visible markers."""
    pattern = re.compile(re.escape(_PENDING) + r"|" + _CITATION_RE.pattern)
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        run = paragraph.add_run(token)
        if token == _PENDING:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.font.color.rgb = RGBColor(0x99, 0x66, 0x00)
        else:
            run.italic = True
            run.font.size = Pt(9)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def write_docx(report_md: str, *, title: str, path: Path | None = None) -> Path:
    parsed = parse_report_markdown(report_md, title=title)
    target = path or default_output_path(parsed.title, "docx")
    target.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(11)

    document.add_heading(parsed.title, level=0)

    section_names = [b.text for b in parsed.blocks if b.kind == "heading" and b.level == 2]
    if section_names:
        _add_toc_paragraph(document, section_names)

    current_section_refs: list[tuple[str, int]] = []

    def flush_section_refs() -> None:
        nonlocal current_section_refs
        if not current_section_refs:
            return
        document.add_paragraph("引用：")
        for file_name, page in current_section_refs:
            document.add_paragraph(f"[{file_name} p{page}]", style="List Bullet")
        current_section_refs = []

    for block in parsed.blocks:
        if block.kind == "heading":
            if block.level == 1 and block.text == parsed.title:
                continue
            if block.level == 2:
                flush_section_refs()
            level = min(max(block.level, 1), 4)
            document.add_heading(block.text, level=level)
            continue
        # paragraph
        for match in _CITATION_RE.finditer(block.text):
            key = (match.group(1).strip(), int(match.group(2)))
            if key not in current_section_refs:
                current_section_refs.append(key)
        para = document.add_paragraph()
        _append_runs_with_citations_and_pending(para, block.text)

    flush_section_refs()

    if parsed.citations:
        document.add_heading("全部引用", level=1)
        for file_name, page in parsed.citations:
            document.add_paragraph(f"[{file_name} p{page}]", style="List Bullet")

    document.save(str(target))
    return target
